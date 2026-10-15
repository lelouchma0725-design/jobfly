"""文檔文本抽取（FR-1b → P2 解析流水線）。

- 數字版 PDF：pypdf 直接抽文本（無需 OCR，GAP 決策）
- DOCX：python-docx 段落 + 表格
- TXT：直接讀取
- 圖片/掃描版 PDF：標記 needs_ocr，由 OCR worker（PaddleOCR，可選依賴）兜底
"""
from __future__ import annotations

import io
import logging

from app.core.r2 import r2_client
from app.config import settings

logger = logging.getLogger(__name__)

MIN_TEXT_LEN = 80  # 文本過短視為掃描件，需 OCR


def download_bytes(storage_path: str) -> bytes:
    resp = r2_client().get_object(Bucket=settings.r2_bucket, Key=storage_path)
    return resp["Body"].read()


def extract_text(storage_path: str, mime: str) -> tuple[str, bool]:
    """返回 (文本, 是否需要OCR)。"""
    data = download_bytes(storage_path)
    if mime == "application/pdf":
        text = _pdf_text(data)
        return text, len(text.strip()) < MIN_TEXT_LEN
    if mime == "text/plain":
        text = data.decode("utf-8", errors="replace")
        return text, False
    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _docx_text(data), False
    if mime in ("image/jpeg", "image/png", "image/webp"):
        return "", True  # 純圖片，走 OCR
    return "", True


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def ocr_image(data: bytes) -> str:
    """PaddleOCR 兜底（可選依賴 app[ocr]，未安裝時拋出明確錯誤）。"""
    try:
        from paddleocr import PaddleOCR
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("未安裝 OCR 依賴：pip install -e '.[ocr]'") from exc
    from PIL import Image

    ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
    image = Image.open(io.BytesIO(data)).convert("RGB")
    result = ocr.ocr(image, cls=True)
    lines = []
    for page in result or []:
        for item in page or []:
            if item and len(item) >= 2 and item[1]:
                lines.append(item[1][0])
    return "\n".join(lines)
